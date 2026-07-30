"""Extraction — files to structured content with provenance.

Understands nothing and judges nothing. Its only job is to read accurately
and to say plainly what it could not read.

Three defects found during testing are handled here, all of which fail
silently if left alone:
  * digit runs reverse in RTL context   -> read by glyph position (docs/01 §10.2)
  * empty table cells collapse the row  -> geometric cells + width check (§7.4)
  * a page can extract fine and render as garbage -> integrity check (§3.1)
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

import fitz
import openpyxl

from . import ocr

from .normalize import ARABIC_INDIC, presentation_form_ratio

BBox = tuple[float, float, float, float]


@dataclass
class Block:
    page: int
    bbox: BBox
    text: str
    section_path: str | None = None
    confidence: float = 1.0
    source: str = "native"            # native | ocr
    numeric_reliable: bool = True     # False -> values here must not be compared


@dataclass
class Table:
    page: int
    headers: list[str]
    rows: list[list[str]]
    row_bboxes: list[BBox] = field(default_factory=list)
    issues: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.issues

    def as_dicts(self) -> list[dict]:
        return [dict(zip(self.headers, r)) for r in self.rows]


@dataclass
class PageHealth:
    page: int
    kind: str                     # digital | scanned
    char_count: int
    render_suspect: bool = False
    reason: str | None = None
    ocr_confidence: float | None = None
    scan_quality: float | None = None
    quarantined_numerals: int = 0


@dataclass
class Document:
    path: Path
    name: str
    doc_type: str = "other"
    type_confidence: float = 0.0
    revision: str | None = None
    language: list[str] = field(default_factory=list)
    blocks: list[Block] = field(default_factory=list)
    tables: list[Table] = field(default_factory=list)
    health: list[PageHealth] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n".join(b.text for b in self.blocks)

    @property
    def render_suspect_pages(self) -> list[int]:
        return [h.page for h in self.health if h.render_suspect]


# ── numbers: read by geometry, never from the extracted string ────────────
def _rebuild_digit_runs(chars: list[dict]) -> str:
    """Digits are always left-to-right under UAX #9 whatever surrounds them,
    so a digit run is read by ascending x. No mode detection, no heuristics,
    and an already-correct number is returned unchanged."""
    out, run = [], []

    def flush():
        if run:
            out.extend(c["c"] for c in sorted(run, key=lambda c: c["bbox"][0]))
            run.clear()

    for ch in chars:
        if ch["c"] in ARABIC_INDIC:
            run.append(ch)
        else:
            flush()
            out.append(ch["c"])
    flush()
    return "".join(out)


def _line_chars(line: dict) -> list[dict]:
    return [c for sp in line["spans"] for c in sp["chars"]]


# ── visual integrity (docs/01 §3.1) ──────────────────────────────────────
INK_PER_CHAR_FLOOR = 4.0e-6


def _check_render(page: fitz.Page, char_count: int) -> tuple[bool, str | None]:
    """Text extracting cleanly is not evidence the page is readable."""
    for font in page.get_fonts(full=False):
        xref, ext, ftype, basefont = font[0], font[1], font[2], font[3]
        if ftype != "Type1" and ext == "n/a" and not basefont.startswith(
                ("Helvetica", "Courier", "Times", "Symbol", "ZapfD")):
            return True, f"FONT_NOT_EMBEDDED:{basefont}"

    if char_count > 100:
        pm = page.get_pixmap(dpi=72, colorspace=fitz.csGRAY)
        ink = sum(1 for v in pm.samples if v < 200) / len(pm.samples)
        if ink / char_count < INK_PER_CHAR_FLOOR:
            return True, f"RENDER_SUSPECT:ink_per_char={ink / char_count:.2e}"
    return False, None


# ── PDF ──────────────────────────────────────────────────────────────────
SECTION_RE = re.compile(r"\bSECTION\s+((?:\d{2}\s+){2}\d{2})\b", re.IGNORECASE)
CLAUSE_RE = re.compile(r"^\s*(\d+\.\d+(?:\.\d+)?)\s")


OCR_DPI = 300


def _ocr_page(page: fitz.Page, pno: int, doc: Document, health: PageHealth) -> None:
    """Read a scanned page. Failure degrades the page, never the whole run."""
    provider = ocr.default_provider()
    if not provider.available():
        health.reason = "OCR_UNAVAILABLE:tesseract not installed"
        return

    grey = page.get_pixmap(dpi=OCR_DPI, colorspace=fitz.csGRAY)
    health.scan_quality = round(ocr.scan_quality(grey), 3)
    try:
        result = provider.recognize(page.get_pixmap(dpi=OCR_DPI).tobytes("png"),
                                    ["ara", "eng"])
    except Exception as exc:
        health.reason = f"OCR_FAILED:{type(exc).__name__}"
        return

    scale = 72.0 / OCR_DPI                       # OCR pixels back to PDF points
    for line in result.lines:
        doc.blocks.append(Block(
            page=pno,
            bbox=tuple(v * scale for v in line.bbox),
            text=line.text,
            confidence=round(line.confidence, 3),
            source="ocr",
            numeric_reliable=line.numeric_reliable))

    # A scan carries no PDF table structure, so find_tables() sees nothing and
    # a door schedule would vanish without a word. Rebuild it from word boxes.
    if (rebuilt := ocr.reconstruct_table(result.words)):
        headers, rows = rebuilt
        table = Table(page=pno, headers=headers, rows=rows)
        bad = [r for r in rows if len(r) != len(headers)]
        if bad:
            table.issues.append(f"ROW_WIDTH_MISMATCH:{len(bad)} rows")
        table.issues.append("SOURCE_OCR — cells recovered from word positions")
        doc.tables.append(table)

    health.char_count = len(result.text)
    health.ocr_confidence = round(result.mean_confidence, 3)
    health.quarantined_numerals = result.quarantined_numerals
    if not result.lines:
        health.reason = "NO_TEXT_EXTRACTED:OCR returned nothing"
    for lang in result.languages:
        code = {"ara": "ar", "eng": "en"}.get(lang, lang)
        if code not in doc.language:
            doc.language.append(code)


def extract_pdf(path: Path, ocr_enabled: bool = True) -> Document:
    doc = Document(path=path, name=path.name)
    pdf = fitz.open(path)
    section = None

    for pno, page in enumerate(pdf, 1):
        raw = page.get_text("rawdict")
        chars_total = 0
        # Grouped per block, not per line: a clause's subject and its
        # measurement routinely land on different lines ("External wall
        # thermal insulation …" / "… thickness of 100 mm"), and splitting
        # them loses the link between the two.
        for blk in raw["blocks"]:
            lines, chars_all = [], []
            for line in blk.get("lines", []):
                chars = _line_chars(line)
                if not chars:
                    continue
                chars_all += chars
                lines.append(_rebuild_digit_runs(chars).strip())
            chars_total += len(chars_all)
            text = " ".join(l for l in lines if l).strip()
            if not text:
                continue
            if (m := SECTION_RE.search(text)):
                section = m.group(1)
            clause = CLAUSE_RE.match(text)
            doc.blocks.append(Block(
                page=pno, text=text,
                bbox=(min(c["bbox"][0] for c in chars_all),
                      min(c["bbox"][1] for c in chars_all),
                      max(c["bbox"][2] for c in chars_all),
                      max(c["bbox"][3] for c in chars_all)),
                section_path=f"{section} / {clause.group(1)}" if section and clause
                else section))

        kind = "digital" if chars_total > 50 else "scanned"
        health = PageHealth(pno, kind, chars_total)
        if kind == "digital":
            health.render_suspect, health.reason = _check_render(page, chars_total)
        elif ocr_enabled:
            _ocr_page(page, pno, doc, health)
        doc.health.append(health)

        for tbl in page.find_tables().tables:
            data = tbl.extract()
            if len(data) < 2:
                continue
            headers = [(h or "").strip() for h in data[0]]
            rows = [[(c or "").strip() for c in r] for r in data[1:]]
            t = Table(page=pno, headers=headers, rows=rows)
            # Structural guard: a collapsed empty cell shifts every value after
            # it and no arithmetic check would notice (docs/01 §7.4).
            bad = [i for i, r in enumerate(rows) if len(r) != len(headers)]
            if bad:
                t.issues.append(f"ROW_WIDTH_MISMATCH:{len(bad)} rows")
            doc.tables.append(t)

    full = doc.text
    if re.search(r"[؀-ۿ]", full) and "ar" not in doc.language:
        doc.language.append("ar")
    if re.search(r"[A-Za-z]", full) and "en" not in doc.language:
        doc.language.append("en")
    if presentation_form_ratio(full) > 0.02:
        doc.notes.append("ARABIC_PRESENTATION_FORMS — visual-order extractor")

    pdf.close()
    return doc


# ── Excel ────────────────────────────────────────────────────────────────
def extract_xlsx(path: Path) -> Document:
    doc = Document(path=path, name=path.name)
    wb = openpyxl.load_workbook(path, data_only=True)
    for ws in wb.worksheets:
        rows = [[("" if c is None else str(c).strip()) for c in r]
                for r in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(r)]
        if not rows:
            continue
        hi = next((i for i, r in enumerate(rows)
                   if sum(bool(c) for c in r) >= 3), 0)
        doc.tables.append(Table(page=1, headers=rows[hi], rows=rows[hi + 1:]))
        for r in rows:
            doc.blocks.append(Block(page=1, bbox=(0, 0, 0, 0), text=" | ".join(r)))
    doc.health.append(PageHealth(1, "digital", len(doc.text)))
    doc.language.append("en")
    return doc


def extract(path: Path, ocr_enabled: bool = True) -> Document:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf(path, ocr_enabled=ocr_enabled)
    if suffix in (".xlsx", ".xlsm", ".xls"):
        return extract_xlsx(path)
    raise ValueError(f"unsupported file type: {suffix}")
